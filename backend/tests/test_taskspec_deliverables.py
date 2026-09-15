import pytest
from pathlib import Path
from openpyxl import load_workbook
from pptx import Presentation

from app.agent.task_spec_builder import build_task_spec
from app.api.v1.chat.router import generate_structured_response


def test_ppt_prompt_1_pump_p204b_reliability_risk_briefing():
    query = (
        "Create a 5-slide PowerPoint titled “Pump P-204B Reliability Risk Briefing”. "
        "Include: 1. Executive summary 2. Evidence and operating thresholds "
        "3. Incident history and risk consequences 4. Corrective actions with owners and due dates "
        "5. Management decisions and sources Generate and attach the PPTX file."
    )
    task_spec = build_task_spec(query, mode="generate_ppt", request_id="req_ppt1_12345678")
    assert task_spec.title == "Pump P-204B Reliability Risk Briefing"
    assert task_spec.slide_count == 5

    agent_result = {"intent": "generate_ppt", "tool": "generate_ppt", "sources": [], "task_spec": task_spec}
    res = generate_structured_response(query, agent_result, "req_ppt1_12345678")

    assert res["task_type"] == "generate_ppt"
    assert res["title"] == "Pump P-204B Reliability Risk Briefing"
    assert res["slide_count"] == 5
    assert len(res["files"]) > 0

    file_info = res["files"][0]
    assert file_info["name"].endswith(".pptx")
    assert ("pump_p204b" in file_info["name"].lower() or "pump_p-204b" in file_info["name"].lower())

    # Reopen presentation and verify slide count
    vault_file = Path("data/vault") / file_info["name"]
    assert vault_file.exists() and vault_file.stat().st_size > 1024
    prs = Presentation(vault_file)
    assert len(prs.slides) == 5


def test_ppt_prompt_2_safety_compliance_recovery_plan():
    query = (
        "Create a 5-slide PowerPoint titled “Safety Compliance Recovery Plan — September 2026”. "
        "Include: 1. Overall compliance position 2. Controls below target "
        "3. Permit-to-work and confined-space gaps 4. Housekeeping closure problem and recovery actions "
        "5. Management actions, owners, due dates, and sources Generate and attach the PPTX file."
    )
    task_spec = build_task_spec(query, mode="generate_ppt", request_id="req_ppt2_12345678")
    assert task_spec.title == "Safety Compliance Recovery Plan — September 2026"
    assert task_spec.slide_count == 5

    agent_result = {"intent": "generate_ppt", "tool": "generate_ppt", "sources": [], "task_spec": task_spec}
    res = generate_structured_response(query, agent_result, "req_ppt2_12345678")

    assert res["task_type"] == "generate_ppt"
    assert res["title"] == "Safety Compliance Recovery Plan — September 2026"
    assert res["slide_count"] == 5
    assert len(res["files"]) > 0

    file_info = res["files"][0]
    assert file_info["name"].endswith(".pptx")
    assert "compliance" in file_info["name"].lower()

    vault_file = Path("data/vault") / file_info["name"]
    assert vault_file.exists() and vault_file.stat().st_size > 1024
    prs = Presentation(vault_file)
    assert len(prs.slides) == 5


def test_ppt_prompt_3_july_august_operations_safety_review():
    query = (
        "Create a 6-slide PowerPoint titled “July–August 2026 Operations Safety Review”. "
        "Include: 1. Executive safety summary 2. Incident and near-miss breakdown "
        "3. Pump P-204B operational risk 4. Compliance scorecard 5. Open corrective actions "
        "6. Leadership decisions required this week with sources Generate and attach the PPTX file."
    )
    task_spec = build_task_spec(query, mode="generate_ppt", request_id="req_ppt3_12345678")
    assert task_spec.title == "July–August 2026 Operations Safety Review"
    assert task_spec.slide_count == 6

    agent_result = {"intent": "generate_ppt", "tool": "generate_ppt", "sources": [], "task_spec": task_spec}
    res = generate_structured_response(query, agent_result, "req_ppt3_12345678")

    assert res["task_type"] == "generate_ppt"
    assert res["title"] == "July–August 2026 Operations Safety Review"
    assert res["slide_count"] == 6
    assert len(res["files"]) > 0

    file_info = res["files"][0]
    assert file_info["name"].endswith(".pptx")
    vault_file = Path("data/vault") / file_info["name"]
    assert vault_file.exists() and vault_file.stat().st_size > 1024
    prs = Presentation(vault_file)
    assert len(prs.slides) == 6


def test_excel_prompt_1_incident_register_and_safety_summary():
    query = (
        "Create an Excel workbook titled “Incident Register and Safety Summary”. "
        "Create these sheets: 1. Incident Register 2. Event Summary 3. Severity Summary "
        "The Incident Register must contain: ID, Date, Area, Event Type, Description, Severity, Immediate Action, Status. "
        "Use all incident records from the uploaded document. Generate and attach the XLSX file."
    )
    task_spec = build_task_spec(query, mode="generate_excel", request_id="req_xls1_12345678")
    assert task_spec.title == "Incident Register and Safety Summary"
    assert task_spec.requested_sheets == ["Incident Register", "Event Summary", "Severity Summary"]

    agent_result = {"intent": "generate_excel", "tool": "generate_excel", "sources": [], "task_spec": task_spec}
    res = generate_structured_response(query, agent_result, "req_xls1_12345678")

    assert res["task_type"] == "generate_excel"
    assert res["title"] == "Incident Register and Safety Summary"
    assert res["requested_sheets"] == ["Incident Register", "Event Summary", "Severity Summary"]
    assert len(res["files"]) > 0

    file_info = res["files"][0]
    assert file_info["name"].endswith(".xlsx")
    vault_file = Path("data/vault") / file_info["name"]
    assert vault_file.exists() and vault_file.stat().st_size > 1024

    wb = load_workbook(vault_file)
    assert wb.sheetnames == ["Incident Register", "Event Summary", "Severity Summary"]


def test_excel_prompt_2_safety_compliance_tracker():
    query = (
        "Create an Excel workbook titled “Safety Compliance Tracker — July–August 2026”. "
        "Create these sheets: 1. Compliance Scorecard 2. Corrective Actions 3. Dashboard "
        "Include: Control Area, Target, Actual, Score, RAG Status, Gap to Target. Generate and attach the XLSX file."
    )
    task_spec = build_task_spec(query, mode="generate_excel", request_id="req_xls2_12345678")
    assert task_spec.title == "Safety Compliance Tracker — July–August 2026"
    assert task_spec.requested_sheets == ["Compliance Scorecard", "Corrective Actions", "Dashboard"]

    agent_result = {"intent": "generate_excel", "tool": "generate_excel", "sources": [], "task_spec": task_spec}
    res = generate_structured_response(query, agent_result, "req_xls2_12345678")

    assert res["task_type"] == "generate_excel"
    assert res["title"] == "Safety Compliance Tracker — July–August 2026"
    assert res["requested_sheets"] == ["Compliance Scorecard", "Corrective Actions", "Dashboard"]

    file_info = res["files"][0]
    vault_file = Path("data/vault") / file_info["name"]
    assert vault_file.exists()
    wb = load_workbook(vault_file)
    assert wb.sheetnames == ["Compliance Scorecard", "Corrective Actions", "Dashboard"]


def test_excel_prompt_3_open_corrective_action_tracker():
    query = (
        "Create an Excel workbook titled “Open Corrective Action Tracker”. "
        "Create these sheets: 1. Open Actions 2. Priority Summary 3. Due-Date Schedule "
        "Include only Open or Planned actions. Exclude closed actions. Sort the schedule by due date. Generate and attach the XLSX file."
    )
    task_spec = build_task_spec(query, mode="generate_excel", request_id="req_xls3_12345678")
    assert task_spec.title == "Open Corrective Action Tracker"
    assert task_spec.requested_sheets == ["Open Actions", "Priority Summary", "Due-Date Schedule"]
    assert "exclude_closed_cap06" in task_spec.requested_filters

    agent_result = {"intent": "generate_excel", "tool": "generate_excel", "sources": [], "task_spec": task_spec}
    res = generate_structured_response(query, agent_result, "req_xls3_12345678")

    assert res["task_type"] == "generate_excel"
    assert res["title"] == "Open Corrective Action Tracker"
    assert res["requested_sheets"] == ["Open Actions", "Priority Summary", "Due-Date Schedule"]

    file_info = res["files"][0]
    vault_file = Path("data/vault") / file_info["name"]
    assert vault_file.exists()
    wb = load_workbook(vault_file)
    assert wb.sheetnames == ["Open Actions", "Priority Summary", "Due-Date Schedule"]

    # Verify CAP-06 (closed) is not present in Open Actions sheet
    ws_open = wb["Open Actions"]
    rows_text = [str(cell.value) for row in ws_open.iter_rows() for cell in row if cell.value]
    assert "CAP-06" not in rows_text


def test_report_prompt_1_executive_operations_safety_report():
    query = (
        "Generate a Word report titled “Executive Operations Safety Report — July–August 2026”. "
        "Include: 1. Executive summary 2. Safety performance headline 3. Top operational risks "
        "4. Compliance gaps 5. Open critical and high-priority actions 6. Management recommendations "
        "7. Sources used Generate and attach the DOCX file."
    )
    task_spec = build_task_spec(query, mode="generate_report", request_id="req_doc1_12345678")
    assert task_spec.title == "Executive Operations Safety Report — July–August 2026"

    agent_result = {"intent": "report_generation", "tool": "report_generation", "sources": [], "task_spec": task_spec}
    res = generate_structured_response(query, agent_result, "req_doc1_12345678")

    assert res["task_type"] == "generate_report"
    assert res["title"] == "Executive Operations Safety Report — July–August 2026"
    assert len(res["files"]) > 0

    file_info = res["files"][0]
    assert file_info["name"].endswith(".docx")
    vault_file = Path("data/vault") / file_info["name"]
    assert vault_file.exists() and vault_file.stat().st_size > 1024


def test_report_prompt_2_pump_p204b_reliability_assessment():
    query = (
        "Generate a Word report titled “Pump P-204B Reliability and Safety Assessment”. "
        "Include: 1. Asset condition summary 2. Documented evidence 3. Operating threshold comparison "
        "4. Incident history 5. Potential consequences 6. Required actions and due dates "
        "7. Escalation criteria 8. Sources used Do not invent equipment specifications, costs, production loss, or maintenance history. Generate and attach the DOCX file."
    )
    task_spec = build_task_spec(query, mode="generate_report", request_id="req_doc2_12345678")
    assert task_spec.title == "Pump P-204B Reliability and Safety Assessment"

    agent_result = {"intent": "report_generation", "tool": "report_generation", "sources": [], "task_spec": task_spec}
    res = generate_structured_response(query, agent_result, "req_doc2_12345678")

    assert res["task_type"] == "generate_report"
    assert res["title"] == "Pump P-204B Reliability and Safety Assessment"
    assert len(res["files"]) > 0

    file_info = res["files"][0]
    assert file_info["name"].endswith(".docx")
    assert ("pump_p204b" in file_info["name"].lower() or "pump_p-204b" in file_info["name"].lower())
    vault_file = Path("data/vault") / file_info["name"]
    assert vault_file.exists() and vault_file.stat().st_size > 1024


def test_report_prompt_3_corrective_action_status_review():
    query = (
        "Generate a Word report titled “Corrective Action Status Review — September 2026”. "
        "Include: 1. Purpose and reporting period 2. Open and planned actions by priority "
        "3. Critical action details 4. High-priority action details 5. Owner and due-date table "
        "6. Risks of delayed closure 7. Management escalation recommendations 8. Sources used "
        "Exclude closed CAP-06 as an outstanding action. Generate and attach the DOCX file."
    )
    task_spec = build_task_spec(query, mode="generate_report", request_id="req_doc3_12345678")
    assert task_spec.title == "Corrective Action Status Review — September 2026"
    assert "exclude_closed_cap06" in task_spec.requested_filters

    agent_result = {"intent": "report_generation", "tool": "report_generation", "sources": [], "task_spec": task_spec}
    res = generate_structured_response(query, agent_result, "req_doc3_12345678")

    assert res["task_type"] == "generate_report"
    assert res["title"] == "Corrective Action Status Review — September 2026"
    assert len(res["files"]) > 0

    file_info = res["files"][0]
    assert file_info["name"].endswith(".docx")
    vault_file = Path("data/vault") / file_info["name"]
    assert vault_file.exists() and vault_file.stat().st_size > 1024
